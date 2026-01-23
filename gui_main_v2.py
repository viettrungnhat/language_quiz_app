"""
GUI VERSION 2.1 - Giao diện đồ họa chatbot kiểm tra ngôn ngữ + Voice Quiz v3
Sử dụng tkinter (built-in Python)
Xử lý Excel trực tiếp mà không cần chuyển JSON
Voice: gTTS (Online TTS chất lượng) + Google Speech Recognition (STT)
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import openpyxl
import random
from quiz_engine import QuizEngine
from scorer import Scorer
from voice_quiz_v3 import VoiceQuizManagerV3
from pathlib import Path
import json
from threading import Thread


class LanguageQuizGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("🎓 Chatbot Kiểm Tra Ngôn Ngữ v2.1")
        self.root.geometry("950x750")
        self.root.resizable(True, True)
        
        # Biến
        self.quiz_engine = None
        self.voice_manager = VoiceQuizManagerV3()
        self.selected_file = None
        self.data = []
        self.current_question_idx = 0
        self.attempt = 1
        self.quiz_results = []
        self.is_listening = False
        
        # Setup style
        self.root.configure(bg="#f0f0f0")
        style = ttk.Style()
        style.theme_use('clam')
        
        # Tạo notebook (tabs)
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Tab 1: Setup
        self.setup_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.setup_tab, text="📋 Chuẩn bị")
        self._create_setup_tab()
        
        # Tab 2: Quiz Thường
        self.quiz_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.quiz_tab, text="🎯 Kiểm tra")
        self._create_quiz_tab()
        
        # Tab 3: Voice Quiz (MỚI)
        self.voice_quiz_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.voice_quiz_tab, text="🎤 Kiểm tra Giọng Nói")
        self._create_voice_quiz_tab()
        
        # Tab 4: Results
        self.results_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.results_tab, text="📊 Kết quả")
        self._create_results_tab()
    
    def _create_setup_tab(self):
        """Tab chuẩn bị - chọn file Excel"""
        
        # Frame chọn file
        file_frame = ttk.LabelFrame(self.setup_tab, text="1️⃣ Chọn File Excel", padding=15)
        file_frame.pack(fill=tk.X, padx=10, pady=10)
        
        ttk.Button(file_frame, text="📁 Chọn File Excel...", 
                  command=self.select_excel_file, width=30).pack(pady=10)
        
        self.file_label = ttk.Label(file_frame, text="Chưa chọn file", 
                                    font=("Arial", 10), foreground="gray")
        self.file_label.pack(pady=5)
        
        # Frame lựa chọn
        options_frame = ttk.LabelFrame(self.setup_tab, text="2️⃣ Cấu hình Kiểm tra", padding=15)
        options_frame.pack(fill=tk.X, padx=10, pady=10)
        
        # Chọn Sheet
        ttk.Label(options_frame, text="Chọn Sheet:").pack(anchor=tk.W)
        self.sheet_combo = ttk.Combobox(options_frame, state="readonly", width=40)
        self.sheet_combo.pack(pady=5)
        
        # Chọn loại kiểm tra
        ttk.Label(options_frame, text="Loại Kiểm tra:").pack(anchor=tk.W)
        self.quiz_type_var = tk.StringVar(value="meaning")
        quiz_types = [
            ("Meaning - Hỏi ý nghĩa", "meaning"),
            ("Example - Dịch ví dụ", "example"),
            ("Vietnamese - Dịch từ Việt", "vietnamese")
        ]
        for text, value in quiz_types:
            ttk.Radiobutton(options_frame, text=text, variable=self.quiz_type_var, 
                           value=value).pack(anchor=tk.W)
        
        # Chọn số câu
        ttk.Label(options_frame, text="Số Câu Kiểm tra:").pack(anchor=tk.W, pady=(10,0))
        self.num_questions_var = tk.IntVar(value=10)
        for num in [5, 10, 15, 20]:
            ttk.Radiobutton(options_frame, text=f"{num} câu", 
                           variable=self.num_questions_var, value=num).pack(anchor=tk.W)
        ttk.Radiobutton(options_frame, text="Tất cả câu", 
                       variable=self.num_questions_var, value=999).pack(anchor=tk.W)
        
        # Nút bắt đầu
        button_frame = ttk.Frame(options_frame)
        button_frame.pack(fill=tk.X, pady=15)
        
        ttk.Button(button_frame, text="▶️ KIỂM TRA THƯỜNG", 
                  command=self.start_quiz, width=25).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="🎤 KIỂM TRA GIỌNG NÓI", 
                  command=self.start_voice_quiz, width=25).pack(side=tk.LEFT, padx=5)
    
    def select_excel_file(self):
        """Chọn file Excel"""
        file_path = filedialog.askopenfilename(
            title="Chọn file Excel",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        
        if file_path:
            self.selected_file = file_path
            self.file_label.config(text=f"✓ {Path(file_path).name}", foreground="green")
            
            # Load sheet names
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                sheet_names = wb.sheetnames
                self.sheet_combo['values'] = sheet_names
                if sheet_names:
                    self.sheet_combo.current(0)
            except Exception as e:
                messagebox.showerror("Lỗi", f"Lỗi khi đọc file:\n{e}")
    
    def _read_excel_data(self, sheet_name):
        """Đọc dữ liệu từ Excel sheet"""
        try:
            wb = openpyxl.load_workbook(self.selected_file, data_only=True)
            ws = wb[sheet_name]
            
            # Kiểm tra header
            header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
            
            if len(header_row) < 4:
                messagebox.showerror(
                    "Lỗi Format", 
                    f"❌ File Excel không đúng format!\n\n"
                    f"Cần 4 cột: Word | Meaning | Example EN | Example VI\n"
                    f"Hiện tại chỉ có: {len(header_row)} cột\n\n"
                    f"Vui lòng dùng template: python create_templates.py"
                )
                return []
            
            data = []
            question_id = 1
            error_rows = []
            
            # Bỏ qua dòng tiêu đề (dòng 1)
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                # Kiểm tra dòng trống
                if not row[0]:
                    continue
                
                # Kiểm tra có đủ 4 cột không
                if len(row) < 4:
                    error_rows.append(row_idx)
                    continue
                
                try:
                    item = {
                        "id": question_id,
                        "word": str(row[0]).strip() if row[0] else "",
                        "meaning": str(row[1]).strip() if row[1] else "",
                        "example_en": str(row[2]).strip() if row[2] else "",
                        "example_vi": str(row[3]).strip() if row[3] else ""
                    }
                    
                    # Kiểm tra dữ liệu hợp lệ
                    if item["word"] and item["meaning"]:
                        data.append(item)
                        question_id += 1
                except Exception as e:
                    error_rows.append(row_idx)
                    continue
            
            # Cảnh báo nếu có dòng lỗi
            if error_rows:
                messagebox.showwarning(
                    "Cảnh báo", 
                    f"⚠️ Bỏ qua {len(error_rows)} dòng lỗi: {error_rows[:5]}"
                )
            
            if not data:
                messagebox.showerror(
                    "Lỗi", 
                    "❌ Không có dữ liệu hợp lệ trong file Excel!\n\n"
                    "Kiểm tra:\n"
                    "1. Hàng 1 là header (Word, Meaning, Example EN, Example VI)\n"
                    "2. Cột A (Word) không được trống\n"
                    "3. Ít nhất 2 cột không trống: Meaning & Word"
                )
                return []
            
            return data
        except Exception as e:
            messagebox.showerror(
                "Lỗi", 
                f"❌ Lỗi khi đọc file Excel:\n{str(e)}\n\n"
                f"Vui lòng kiểm tra:\n"
                f"1. File có đúng format .xlsx không?\n"
                f"2. Header (hàng 1) có 4 cột không?\n"
                f"3. Dữ liệu bắt đầu từ hàng 2?"
            )
            return []
    
    def _create_quiz_tab(self):
        """Tab kiểm tra thường (không giọng nói)"""
        
        # Frame câu hỏi
        question_frame = ttk.LabelFrame(self.quiz_tab, text="❓ Câu hỏi", padding=10)
        question_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        self.progress_label = ttk.Label(question_frame, text="")
        self.progress_label.pack(anchor=tk.W)
        
        self.progress_bar = ttk.Progressbar(question_frame, length=400, mode='determinate')
        self.progress_bar.pack(fill=tk.X, pady=5)
        
        self.question_text = scrolledtext.ScrolledText(
            question_frame, height=6, width=60, font=("Arial", 12), wrap=tk.WORD
        )
        self.question_text.pack(fill=tk.BOTH, expand=True, pady=10)
        self.question_text.config(state=tk.DISABLED)
        
        # Frame trả lời
        answer_frame = ttk.LabelFrame(self.quiz_tab, text="💬 Câu trả lời", padding=10)
        answer_frame.pack(fill=tk.X, padx=10, pady=10)
        
        ttk.Label(answer_frame, text="Nhập câu trả lời:").pack(anchor=tk.W)
        self.answer_entry = ttk.Entry(answer_frame, width=60, font=("Arial", 11))
        self.answer_entry.pack(fill=tk.X, pady=5)
        self.answer_entry.bind('<Return>', lambda e: self.submit_answer())
        
        # Frame nút bấm
        button_frame = ttk.Frame(answer_frame)
        button_frame.pack(fill=tk.X, pady=10)
        
        ttk.Button(button_frame, text="💡 Gợi ý", 
                  command=self.show_hint, width=15).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="👁️ Xem đáp án", 
                  command=self.show_answer, width=15).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="✅ Nộp bài", 
                  command=self.submit_answer, width=15).pack(side=tk.LEFT, padx=5)
        
        # Frame phản hồi
        feedback_frame = ttk.LabelFrame(self.quiz_tab, text="📝 Phản hồi", padding=10)
        feedback_frame.pack(fill=tk.X, padx=10, pady=10)
        
        self.feedback_text = scrolledtext.ScrolledText(
            feedback_frame, height=4, width=60, font=("Arial", 10), wrap=tk.WORD
        )
        self.feedback_text.pack(fill=tk.BOTH, expand=True)
        self.feedback_text.config(state=tk.DISABLED)
    
    def _create_voice_quiz_tab(self):
        """Tab kiểm tra bằng giọng nói (MỚI)"""
        
        # Frame thông tin
        info_frame = ttk.LabelFrame(self.voice_quiz_tab, text="ℹ️ Hướng dẫn", padding=10)
        info_frame.pack(fill=tk.X, padx=10, pady=10)
        
        info_text = """
🎤 KIỂM TRA BẰNG GIỌNG NÓI:
1. Máy sẽ phát câu hỏi bằng giọng nói
2. Bạn nói câu trả lời vào microphone
3. Máy sẽ nhận dạng và đánh giá câu trả lời
4. Tự động chuyển sang câu hỏi tiếp theo

⚠️ Yêu cầu: Microphone phải được kết nối
        """
        ttk.Label(info_frame, text=info_text, justify=tk.LEFT, font=("Arial", 9)).pack()
        
        # Frame câu hỏi
        question_frame = ttk.LabelFrame(self.voice_quiz_tab, text="❓ Câu hỏi", padding=10)
        question_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        self.voice_progress_label = ttk.Label(question_frame, text="")
        self.voice_progress_label.pack(anchor=tk.W)
        
        self.voice_progress_bar = ttk.Progressbar(question_frame, length=400, mode='determinate')
        self.voice_progress_bar.pack(fill=tk.X, pady=5)
        
        self.voice_question_text = scrolledtext.ScrolledText(
            question_frame, height=8, width=60, font=("Arial", 12, "bold"), wrap=tk.WORD
        )
        self.voice_question_text.pack(fill=tk.BOTH, expand=True, pady=10)
        self.voice_question_text.config(state=tk.DISABLED)
        
        # Frame trả lời bằng giọng nói
        answer_frame = ttk.LabelFrame(self.voice_quiz_tab, text="🎤 Câu trả lời", padding=10)
        answer_frame.pack(fill=tk.X, padx=10, pady=10)
        
        ttk.Label(answer_frame, text="Bạn nói (nhận dạng bằng giọng nói):", font=("Arial", 10)).pack(anchor=tk.W)
        self.voice_answer_text = scrolledtext.ScrolledText(
            answer_frame, height=3, width=60, font=("Arial", 10), wrap=tk.WORD
        )
        self.voice_answer_text.pack(fill=tk.X, pady=5)
        self.voice_answer_text.config(state=tk.DISABLED)
        
        # Frame nút bấm
        button_frame = ttk.Frame(answer_frame)
        button_frame.pack(fill=tk.X, pady=10)
        
        ttk.Label(button_frame, text="⏳ Tự động lắng nghe từ microphone...", 
                 font=("Arial", 10, "bold"), foreground="blue").pack(side=tk.LEFT, padx=5)
        
        ttk.Button(button_frame, text="⏭️ CÂUHỎI TIẾP", 
                  command=self.voice_next_question, width=20).pack(side=tk.LEFT, padx=5)
        
        ttk.Button(button_frame, text="❌ DỪNG KIỂM TRA", 
                  command=self.voice_stop_quiz, width=20).pack(side=tk.LEFT, padx=5)
        
        # Frame phản hồi
        feedback_frame = ttk.LabelFrame(self.voice_quiz_tab, text="📝 Phản hồi", padding=10)
        feedback_frame.pack(fill=tk.X, padx=10, pady=10)
        
        self.voice_feedback_text = scrolledtext.ScrolledText(
            feedback_frame, height=3, width=60, font=("Arial", 10), wrap=tk.WORD
        )
        self.voice_feedback_text.pack(fill=tk.BOTH, expand=True)
        self.voice_feedback_text.config(state=tk.DISABLED)
    
    def _create_results_tab(self):
        """Tab hiển thị kết quả"""
        
        button_frame = ttk.Frame(self.results_tab)
        button_frame.pack(fill=tk.X, padx=10, pady=10)
        
        ttk.Button(button_frame, text="💾 Lưu kết quả", 
                  command=self.save_results).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="🔄 Kiểm tra lại", 
                  command=self.restart_quiz).pack(side=tk.LEFT, padx=5)
        
        self.results_text = scrolledtext.ScrolledText(
            self.results_tab, font=("Arial", 10), wrap=tk.WORD
        )
        self.results_text.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        self.results_text.config(state=tk.DISABLED)
    
    # ===== QUIZ THƯỜNG =====
    
    def start_quiz(self):
        """Bắt đầu kiểm tra thường"""
        if not self.selected_file:
            messagebox.showerror("Lỗi", "❌ Vui lòng chọn file Excel trước!")
            return
        
        sheet_name = self.sheet_combo.get()
        if not sheet_name:
            messagebox.showerror("Lỗi", "❌ Vui lòng chọn sheet!")
            return
        
        self.data = self._read_excel_data(sheet_name)
        if not self.data:
            return
        
        # Tạo quiz engine
        self.quiz_type_str = self.quiz_type_var.get()
        num_questions = self.num_questions_var.get()
        if num_questions == 999:
            num_questions = len(self.data)
        else:
            num_questions = min(num_questions, len(self.data))
        
        # Khởi tạo QuizEngine với questions và language
        self.quiz_engine = QuizEngine(self.data, language="English")
        # Trộn câu hỏi theo số lượng
        self.quiz_engine.shuffle_questions(num_questions)
        # Lưu loại kiểm tra
        self.quiz_engine.quiz_type = self.quiz_type_str
        
        self.current_question_idx = 0
        self.attempt = 1
        self.quiz_results = []
        
        # Chuyển sang tab kiểm tra
        self.notebook.select(1)
        self.display_question()
    
    def display_question(self):
        """Hiển thị câu hỏi"""
        if self.current_question_idx >= len(self.quiz_engine.questions):
            self.show_results()
            return
        
        question = self.quiz_engine.questions[self.current_question_idx]
        question_text = self.quiz_engine.format_question(question, self.quiz_type_str)
        
        self.question_text.config(state=tk.NORMAL)
        self.question_text.delete(1.0, tk.END)
        self.question_text.insert(tk.END, question_text)
        self.question_text.config(state=tk.DISABLED)
        
        # Cập nhật progress
        current = self.current_question_idx + 1
        total = len(self.quiz_engine.questions)
        self.progress_label.config(
            text=f"Câu {current}/{total} (Lần {self.attempt}/3)"
        )
        self.progress_bar['value'] = (current / total) * 100
        
        # Clear fields
        self.answer_entry.delete(0, tk.END)
        self.feedback_text.config(state=tk.NORMAL)
        self.feedback_text.delete(1.0, tk.END)
        self.feedback_text.config(state=tk.DISABLED)
        
        self.answer_entry.focus()
    
    def submit_answer(self):
        """Nộp câu trả lời"""
        user_answer = self.answer_entry.get().strip()
        
        if not user_answer:
            messagebox.showwarning("Cảnh báo", "❌ Vui lòng nhập câu trả lời!")
            return
        
        question = self.quiz_engine.questions[self.current_question_idx]
        correct_answer = question.get("meaning") if self.quiz_type_str == "meaning" else \
                        question.get("example_en") if self.quiz_type_str == "example" else \
                        question.get("example_vi")
        
        is_correct, feedback = self.quiz_engine.check_answer(
            user_answer, correct_answer, self.attempt
        )
        
        score = 10 if is_correct and self.attempt == 1 else \
               7 if is_correct and self.attempt == 2 else \
               4 if is_correct else 0
        
        # Lưu kết quả
        self.quiz_results.append({
            "question": question.get("word"),
            "user_answer": user_answer,
            "correct_answer": correct_answer,
            "score": score,
            "attempt": self.attempt
        })
        
        # Hiển thị phản hồi
        self.feedback_text.config(state=tk.NORMAL)
        self.feedback_text.delete(1.0, tk.END)
        self.feedback_text.insert(tk.END, f"{feedback}\n\n💡 Đáp án đúng: {correct_answer}")
        self.feedback_text.config(state=tk.DISABLED)
        
        if is_correct or self.attempt >= 3:
            self.root.after(2000, self.next_question)
        else:
            self.attempt += 1
            self.display_question()
    
    def next_question(self):
        """Chuyển sang câu hỏi tiếp theo"""
        self.current_question_idx += 1
        self.attempt = 1
        self.display_question()
    
    def show_hint(self):
        """Hiển thị gợi ý"""
        question = self.quiz_engine.questions[self.current_question_idx]
        hint = self.quiz_engine.get_hint(question)
        
        self.feedback_text.config(state=tk.NORMAL)
        self.feedback_text.delete(1.0, tk.END)
        self.feedback_text.insert(tk.END, f"💡 Gợi ý: {hint}")
        self.feedback_text.config(state=tk.DISABLED)
    
    def show_answer(self):
        """Hiển thị đáp án (0 điểm)"""
        question = self.quiz_engine.questions[self.current_question_idx]
        correct_answer = question.get("meaning") if self.quiz_type_str == "meaning" else \
                        question.get("example_en") if self.quiz_type_str == "example" else \
                        question.get("example_vi")
        
        # Lưu kết quả 0 điểm
        self.quiz_results.append({
            "question": question.get("word"),
            "user_answer": "(Xem đáp án)",
            "correct_answer": correct_answer,
            "score": 0,
            "attempt": self.attempt
        })
        
        self.feedback_text.config(state=tk.NORMAL)
        self.feedback_text.delete(1.0, tk.END)
        self.feedback_text.insert(tk.END, f"👁️ Đáp án: {correct_answer}\n\n(0 điểm)")
        self.feedback_text.config(state=tk.DISABLED)
        
        self.root.after(2000, self.next_question)
    
    # ===== VOICE QUIZ (MỚI) =====
    
    def start_voice_quiz(self):
        """Bắt đầu kiểm tra bằng giọng nói"""
        if not self.selected_file:
            messagebox.showerror("Lỗi", "❌ Vui lòng chọn file Excel trước!")
            return
        
        sheet_name = self.sheet_combo.get()
        if not sheet_name:
            messagebox.showerror("Lỗi", "❌ Vui lòng chọn sheet!")
            return
        
        self.data = self._read_excel_data(sheet_name)
        if not self.data:
            return
        
        # Tạo quiz engine
        self.quiz_type_str = self.quiz_type_var.get()
        num_questions = self.num_questions_var.get()
        if num_questions == 999:
            num_questions = len(self.data)
        else:
            num_questions = min(num_questions, len(self.data))
        
        # Khởi tạo QuizEngine với questions và language
        self.quiz_engine = QuizEngine(self.data, language="English")
        # Trộn câu hỏi theo số lượng
        self.quiz_engine.shuffle_questions(num_questions)
        # Lưu loại kiểm tra
        self.quiz_engine.quiz_type = self.quiz_type_str
        
        self.current_question_idx = 0
        self.quiz_results = []
        
        # Chuyển sang tab voice quiz
        self.notebook.select(2)
        
        # Bắt đầu với câu hỏi đầu tiên
        messagebox.showinfo("Bắt đầu", "🎤 Kiểm tra bằng giọng nói sẽ bắt đầu ngay!\n\nĐảm bảo microphone đã kết nối.")
        self.voice_display_question()
    
    def voice_display_question(self):
        """Hiển thị câu hỏi giọng nói"""
        if self.current_question_idx >= len(self.quiz_engine.questions):
            self.voice_show_results()
            return
        
        question = self.quiz_engine.questions[self.current_question_idx]
        question_text = self.quiz_engine.format_question(question, self.quiz_type_str)
        
        # Hiển thị câu hỏi
        self.voice_question_text.config(state=tk.NORMAL)
        self.voice_question_text.delete(1.0, tk.END)
        self.voice_question_text.insert(tk.END, f"❓ {question_text}")
        self.voice_question_text.config(state=tk.DISABLED)
        
        # Cập nhật progress
        current = self.current_question_idx + 1
        total = len(self.quiz_engine.questions)
        self.voice_progress_label.config(text=f"Câu {current}/{total}")
        self.voice_progress_bar['value'] = (current / total) * 100
        
        # Clear phần trả lời
        self.voice_answer_text.config(state=tk.NORMAL)
        self.voice_answer_text.delete(1.0, tk.END)
        self.voice_answer_text.config(state=tk.DISABLED)
        
        self.voice_feedback_text.config(state=tk.NORMAL)
        self.voice_feedback_text.delete(1.0, tk.END)
        self.voice_feedback_text.config(state=tk.DISABLED)
        
        # Phát câu hỏi bằng giọng nói (chạy trên thread khác)
        Thread(target=self._speak_question, args=(question_text,), daemon=True).start()
    
    def _speak_question(self, question_text):
        """
        Phát câu hỏi bằng giọng nói 2 lần + đếm ngược 3s + tự động lắng nghe + feedback (thread)
        """
        try:
            import time
            
            # 1. Đọc câu hỏi 2 lần + đếm ngược + lắng nghe
            user_answer = self.voice_manager.ask_question_and_listen(
                question_text=question_text,
                language_tts="en"
            )
            
            if not user_answer:
                self.root.after(0, lambda: self.voice_feedback_text.config(state=tk.NORMAL))
                self.root.after(0, lambda: self.voice_feedback_text.delete(1.0, tk.END))
                self.root.after(0, lambda: self.voice_feedback_text.insert(tk.END, "❌ Không nhận dạng được giọng nói"))
                self.root.after(0, lambda: self.voice_feedback_text.config(state=tk.DISABLED))
                return
            
            # 2. Hiển thị câu bạn nói
            self.root.after(0, lambda: self.voice_answer_text.config(state=tk.NORMAL))
            self.root.after(0, lambda: self.voice_answer_text.delete(1.0, tk.END))
            self.root.after(0, lambda: self.voice_answer_text.insert(tk.END, f"🎤 Bạn nói: {user_answer}"))
            self.root.after(0, lambda: self.voice_answer_text.config(state=tk.DISABLED))
            
            # 3. So sánh câu trả lời
            question = self.quiz_engine.questions[self.current_question_idx]
            correct_answer = question.get("meaning") if self.quiz_type_str == "meaning" else \
                            question.get("example_en") if self.quiz_type_str == "example" else \
                            question.get("example_vi")
            
            is_correct, similarity, _ = self.voice_manager.voice_manager.compare_answers(user_answer, correct_answer)
            
            # 4. Lưu kết quả
            score = min(10, int(similarity * 10)) if is_correct else max(0, int(similarity * 5))
            self.quiz_results.append({
                "question": question.get("word"),
                "user_answer": user_answer,
                "correct_answer": correct_answer,
                "score": score,
                "attempt": 1
            })
            
            # 5. Phát feedback bằng giọng nói + hiển thị
            time.sleep(1)
            self.voice_manager.give_feedback(is_correct, user_answer, correct_answer, similarity)
            
            # 6. Tự động chuyển sang câu tiếp theo
            time.sleep(2)
            self.root.after(0, self.voice_next_question)
        
        except Exception as e:
            print(f"❌ Lỗi: {e}")
            error_msg = str(e)
            self.root.after(0, lambda: messagebox.showerror("Lỗi", f"Lỗi: {error_msg}"))
    
    def voice_next_question(self):
        """Chuyển sang câu hỏi tiếp theo (voice quiz)"""
        self.current_question_idx += 1
        self.voice_display_question()
    
    def voice_stop_quiz(self):
        """Dừng kiểm tra giọng nói"""
        if messagebox.askyesno("Xác nhận", "❌ Bạn chắc chắn muốn dừng kiểm tra?"):
            self.voice_show_results()
    
    def voice_show_results(self):
        """Hiển thị kết quả voice quiz"""
        self.show_results()
        self.notebook.select(3)
    
    # ===== RESULTS =====
    
    def show_results(self):
        """Hiển thị kết quả"""
        if not self.quiz_results:
            return
        
        # Tính điểm từ quiz_results
        total_points = sum(r.get("score", 0) for r in self.quiz_results)
        num_questions = len(self.quiz_results)
        avg_score = (total_points / (num_questions * 10)) * 100 if num_questions > 0 else 0
        
        # Xếp loại
        if avg_score >= 90:
            grade = "A - Xuất sắc"
        elif avg_score >= 80:
            grade = "B - Tốt"
        elif avg_score >= 70:
            grade = "C - Khá"
        elif avg_score >= 60:
            grade = "D - Đạt"
        else:
            grade = "F - Chưa đạt"
        
        # Tạo báo cáo
        report = f"""
╔══════════════════════════════════════════════════════════════╗
║               📊 KẾT QUẢ KIỂM TRA                           ║
╚══════════════════════════════════════════════════════════════╝

📈 ĐIỂM TỔNG HỢP:
   Điểm trung bình: {avg_score:.1f}/100
   Xếp loại: {grade}

📋 CHI TIẾT TỪNG CÂU:
"""
        for i, result in enumerate(self.quiz_results, 1):
            report += f"\n{i}. {result['question']} (Lần {result['attempt']})\n"
            report += f"   Bạn trả lời: {result['user_answer']}\n"
            report += f"   Đáp án: {result['correct_answer']}\n"
            report += f"   Điểm: {result['score']}/10\n"
        
        # Tính thống kê
        correct_1st = sum(1 for r in self.quiz_results if r.get("attempt") == 1 and r.get("score") > 0)
        correct_2nd = sum(1 for r in self.quiz_results if r.get("attempt") == 2 and r.get("score") > 0)
        correct_3rd = sum(1 for r in self.quiz_results if r.get("attempt") == 3 and r.get("score") > 0)
        total_wrong = sum(1 for r in self.quiz_results if r.get("score") == 0)
        
        report += f"""
📊 THỐNG KÊ:
   Tổng câu hỏi: {num_questions}
   Đúng lần 1: {correct_1st}
   Đúng lần 2: {correct_2nd}
   Đúng lần 3: {correct_3rd}
   Sai: {total_wrong}
"""
        
        self.results_text.config(state=tk.NORMAL)
        self.results_text.delete(1.0, tk.END)
        self.results_text.insert(tk.END, report)
        self.results_text.config(state=tk.DISABLED)
    
    def save_results(self):
        """Lưu kết quả"""
        if not self.quiz_results:
            messagebox.showwarning("Cảnh báo", "❌ Không có kết quả để lưu!")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".json",
            filetypes=[("JSON files", "*.json"), ("All files", "*.*")]
        )
        
        if file_path:
            try:
                with open(file_path, 'w', encoding='utf-8') as f:
                    json.dump(self.quiz_results, f, ensure_ascii=False, indent=2)
                messagebox.showinfo("Thành công", f"✅ Kết quả đã lưu:\n{file_path}")
            except Exception as e:
                messagebox.showerror("Lỗi", f"❌ Lỗi khi lưu:\n{e}")
    
    def restart_quiz(self):
        """Kiểm tra lại"""
        self.notebook.select(0)


def main():
    root = tk.Tk()
    app = LanguageQuizGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()
